import os
import datetime
import uuid
import json
import io
from typing import List, Optional, Dict, Any

import uvicorn
from fastapi import FastAPI, Depends, HTTPException, status, Request, Query, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from fastapi.security import OAuth2PasswordBearer
from fastapi.responses import FileResponse, JSONResponse, HTMLResponse
from fastapi.staticfiles import StaticFiles

from pydantic import BaseModel, Field

from jose import JWTError, jwt

from dotenv import load_dotenv
from google.oauth2 import id_token as google_id_token
from google.auth.transport import requests as google_requests
from google.oauth2.credentials import Credentials
from google_auth_oauthlib.flow import Flow
from googleapiclient.discovery import build
import google.generativeai as genai

import firebase_admin
from firebase_admin import credentials, firestore

import openpyxl
import docx
from pypdf import PdfReader

import traceback

load_dotenv()

GOOGLE_CLIENT_ID = os.getenv("GOOGLE_CLIENT_ID")
JWT_SECRET_KEY = os.getenv("JWT_SECRET_KEY")
JWT_ALGORITHM = os.getenv("JWT_ALGORITHM")
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")
FIREBASE_SERVICE_ACCOUNT_KEY_PATH = os.getenv("FIREBASE_SERVICE_ACCOUNT_KEY_PATH")

CLIENT_SECRET_FILE = 'client_secret.json'
SCOPES = ['https://www.googleapis.com/auth/calendar.events']
REDIRECT_URI = "http://127.0.0.1:8000/auth/google/calendar/callback"

if GEMINI_API_KEY:
    genai.configure(api_key=GEMINI_API_KEY)
else:
    print("Warning: GEMINI_API_KEY not found. AI features will be disabled.")

if not FIREBASE_SERVICE_ACCOUNT_KEY_PATH or not os.path.exists(FIREBASE_SERVICE_ACCOUNT_KEY_PATH):
    raise ValueError("Firebase service account key path not found or invalid.")
cred = credentials.Certificate(FIREBASE_SERVICE_ACCOUNT_KEY_PATH)
firebase_admin.initialize_app(cred)
db = firestore.client()

app = FastAPI(title="SnapTask API with Firebase")

@app.middleware("http")
async def add_security_headers(request: Request, call_next):
    response = await call_next(request)
    response.headers["Cross-Origin-Opener-Policy"] = "same-origin-allow-popups"
    response.headers["Cross-Origin-Embedder-Policy"] = "require-corp"
    return response

app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:5500",
        "http://127.0.0.1:5500",
        "http://localhost:8000",
        "http://127.0.0.1:8000"
    ],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

oauth2_scheme = OAuth2PasswordBearer(tokenUrl="/auth/google/callback")

class TaskBase(BaseModel):
    description: str
    priority: str = "medium"
    estimated_duration_minutes: int = 30
    start_time: Optional[str] = None
    end_time: Optional[str] = None
    is_daily_routine: bool = False

class TaskCreate(TaskBase):
    pass

class TaskSchema(TaskBase):
    id: str
    status: str
    owner_email: str

class Availability(BaseModel):
    date: str
    start_time: str
    end_time: str

class ScheduleRequest(BaseModel):
    tasks: List[TaskCreate]
    availability: List[Availability]

class ScheduledTask(BaseModel):
    task_description: str
    start_time: str
    end_time: str
    priority: str
    is_daily_routine: bool

class AIScheduleResponse(BaseModel):
    schedule_id: str
    suggested_schedule: List[ScheduledTask]
    notes: str

class CalendarSyncRequest(BaseModel):
    schedule: List[ScheduledTask]

class OcrResponse(BaseModel):
    tasks: List[str]

def create_access_token(data: dict, expires_delta: Optional[datetime.timedelta] = None):
    to_encode = data.copy()
    expire = datetime.datetime.now(datetime.timezone.utc) + (expires_delta or datetime.timedelta(days=1))
    to_encode.update({"exp": expire})
    return jwt.encode(to_encode, JWT_SECRET_KEY, algorithm=JWT_ALGORITHM)

async def get_current_user(token: str = Depends(oauth2_scheme)) -> Dict[str, Any]:
    credentials_exception = HTTPException(
        status_code=status.HTTP_401_UNAUTHORIZED,
        detail="Could not validate credentials",
        headers={"WWW-Authenticate": "Bearer"},
    )
    try:
        payload = jwt.decode(token, JWT_SECRET_KEY, algorithms=[JWT_ALGORITHM])
        email: str = payload.get("email")
        if email is None:
            raise credentials_exception
    except JWTError:
        raise credentials_exception
    
    user_doc_ref = db.collection('users').document(email)
    user_doc = user_doc_ref.get()
    if not user_doc.exists:
        raise credentials_exception
    
    user_data = user_doc.to_dict()
    user_data['email'] = email
    return user_data

@app.post("/auth/google/callback", tags=["Authentication"])
async def google_auth_callback(google_token: dict):
    try:
        if not GOOGLE_CLIENT_ID:
            raise HTTPException(status_code=500, detail="GOOGLE_CLIENT_ID is not set on the backend.")
        idinfo = google_id_token.verify_oauth2_token(google_token['token'], google_requests.Request(), GOOGLE_CLIENT_ID)
        email = idinfo['email']
        user_doc_ref = db.collection('users').document(email)
        user_doc = user_doc_ref.get()
        if not user_doc.exists:
            user_data = {
                "email": email,
                "name": idinfo.get('name', 'New User'),
                "picture_url": idinfo.get('picture'),
                "created_at": firestore.SERVER_TIMESTAMP,
                "calendar_credentials": None
            }
            user_doc_ref.set(user_data)
        app_token = create_access_token(data={"email": email})
        return {
            "app_token": app_token,
            "user_info": {
                "name": idinfo.get('name'),
                "email": email,
                "picture_url": idinfo.get('picture')
            }
        }
    except ValueError as e:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail=f"Invalid Google token: {e}")
    except Exception as e:
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=f"Internal error: {e}")

@app.get("/auth/google/calendar/login", tags=["Google Calendar"])
def calendar_auth_login(current_user: Dict[str, Any] = Depends(get_current_user)):
    user_email = current_user['email']
    if not os.path.exists(CLIENT_SECRET_FILE):
        raise HTTPException(status_code=500, detail="Google Client Secret file not found on server.")
    
    flow = Flow.from_client_secrets_file(
        CLIENT_SECRET_FILE,
        scopes=SCOPES,
        redirect_uri=REDIRECT_URI
    )
    authorization_url, state = flow.authorization_url(
        access_type='offline',
        include_granted_scopes='true',
        state=user_email
    )
    return {"authorization_url": authorization_url}

@app.get("/auth/google/calendar/callback", tags=["Google Calendar"], response_class=HTMLResponse)
async def calendar_auth_callback(request: Request, code: str = Query(...), state: str = Query(...)):
    user_email = state
    if not os.path.exists(CLIENT_SECRET_FILE):
        raise HTTPException(status_code=500, detail="Google Client Secret file not found on server.")
    
    flow = Flow.from_client_secrets_file(
        CLIENT_SECRET_FILE,
        scopes=SCOPES,
        redirect_uri=REDIRECT_URI
    )
    flow.fetch_token(code=code)
    
    creds = flow.credentials
    creds_json = {
        'token': creds.token,
        'refresh_token': creds.refresh_token,
        'token_uri': creds.token_uri,
        'client_id': creds.client_id,
        'client_secret': creds.client_secret,
        'scopes': creds.scopes
    }

    user_doc_ref = db.collection('users').document(user_email)
    user_doc_ref.update({"calendar_credentials": json.dumps(creds_json)})
    
    return """
    <html><head><title>Authentication Successful</title></head><body><p>Authentication successful! You can close this window now.</p><script>window.close();</script></body></html>
    """

@app.get("/api/v1/dashboard/summary", tags=["Dashboard"])
async def get_dashboard_summary(current_user: Dict[str, Any] = Depends(get_current_user)):
    user_email = current_user['email']
    tasks_ref = db.collection('tasks').where('owner_email', '==', user_email).stream()
    total_tasks = 0
    completed_tasks = 0
    for task in tasks_ref:
        total_tasks += 1
        if task.to_dict().get('status') == 'completed':
            completed_tasks += 1
    return {"totalTasks": total_tasks, "completedTasks": completed_tasks}

@app.get("/tasks/", response_model=List[TaskSchema], tags=["Tasks"])
async def read_tasks(current_user: Dict[str, Any] = Depends(get_current_user)):
    user_email = current_user['email']
    tasks_stream = db.collection('tasks').where('owner_email', '==', user_email).stream()
    tasks_list = []
    for task in tasks_stream:
        task_data = task.to_dict()
        task_data['id'] = task.id
        tasks_list.append(task_data)
    return tasks_list

@app.patch("/tasks/{task_id}/complete", status_code=status.HTTP_200_OK, tags=["Tasks"])
async def complete_task(task_id: str, current_user: Dict[str, Any] = Depends(get_current_user)):
    user_email = current_user['email']
    task_ref = db.collection('tasks').document(task_id)
    task_doc = task_ref.get()

    if not task_doc.exists:
        raise HTTPException(status_code=404, detail="Task not found")
    
    if task_doc.to_dict().get('owner_email') != user_email:
        raise HTTPException(status_code=403, detail="Not authorized to modify this task")

    task_ref.update({"status": "completed"})
    return {"message": "Task marked as completed."}

@app.delete("/tasks/{task_id}", status_code=status.HTTP_200_OK, tags=["Tasks"])
async def delete_task(task_id: str, current_user: Dict[str, Any] = Depends(get_current_user)):
    user_email = current_user['email']
    task_ref = db.collection('tasks').document(task_id)
    task_doc = task_ref.get()

    if not task_doc.exists:
        raise HTTPException(status_code=404, detail="Task not found")

    if task_doc.to_dict().get('owner_email') != user_email:
        raise HTTPException(status_code=403, detail="Not authorized to delete this task")

    task_ref.delete()
    return {"message": "Task deleted successfully."}

@app.delete("/tasks/reset", status_code=status.HTTP_204_NO_CONTENT, tags=["Tasks"])
async def reset_tasks(current_user: Dict[str, Any] = Depends(get_current_user)):
    user_email = current_user['email']
    tasks_query = db.collection('tasks').where('owner_email', '==', user_email)
    tasks_docs = tasks_query.stream()

    batch = db.batch()
    doc_count = 0
    for doc in tasks_docs:
        batch.delete(doc.reference)
        doc_count += 1
    
    if doc_count > 0:
        batch.commit()
    
    return {"message": f"{doc_count} tasks have been reset."}

@app.post("/process/file/extract_tasks", tags=["AI Processing"])
async def extract_tasks_from_file(
    file: UploadFile = File(...),
    current_user: Dict[str, Any] = Depends(get_current_user)
):
    if not GEMINI_API_KEY or not genai:
        raise HTTPException(status_code=503, detail="AI Service is not configured or available.")

    content_type = file.content_type
    file_content = await file.read()
    extracted_text = ""

    try:
        if "image" in content_type:
            model = genai.GenerativeModel('gemini-1.5-flash')
            image_parts = [{"mime_type": content_type, "data": file_content}]
            prompt = "Analyze the image and extract all distinct tasks or to-do list items. Return a JSON object with a 'tasks' key containing an array of strings."
            response = await model.generate_content_async([prompt, *image_parts])
            extracted_text = response.text

        elif content_type == 'application/pdf':
            reader = PdfReader(io.BytesIO(file_content))
            for page in reader.pages:
                extracted_text += page.extract_text() + "\n"
        
        elif content_type == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document':
            doc = docx.Document(io.BytesIO(file_content))
            for para in doc.paragraphs:
                extracted_text += para.text + "\n"
        
        elif content_type == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet':
            workbook = openpyxl.load_workbook(io.BytesIO(file_content))
            sheet = workbook.active
            for row in sheet.iter_rows(values_only=True):
                extracted_text += " ".join([str(cell) for cell in row if cell is not None]) + "\n"
        
        else:
            raise HTTPException(status_code=400, detail=f"Unsupported file type: {content_type}")

        if "image" not in content_type:
            model = genai.GenerativeModel('gemini-1.5-flash')
            prompt = f"""
            Analyze the following text content extracted from a document.
            Identify all distinct tasks or to-do list items from this text.
            Return the result as a single, valid JSON object with one key: "tasks".
            The value of "tasks" should be an array of strings, where each string is a separate task you identified.
            If no tasks are found, return an empty array.
            Here is the text:
            ---
            {extracted_text}
            ---
            """
            response = await model.generate_content_async(prompt)
            cleaned_response_text = response.text.strip().replace("```json", "").replace("```", "")
        else:
             cleaned_response_text = extracted_text.strip().replace("```json", "").replace("```", "")

        ai_data = json.loads(cleaned_response_text)
        return {"tasks": ai_data.get("tasks", [])}

    except Exception as e:
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=f"Failed to process file: {str(e)}")

@app.post("/process/ai/generate_schedule", response_model=AIScheduleResponse, tags=["AI Processing"])
async def generate_ai_schedule(schedule_request: ScheduleRequest, current_user: Dict[str, Any] = Depends(get_current_user)):
    if not GEMINI_API_KEY or not genai:
        raise HTTPException(status_code=503, detail="AI Service is not configured or available.")
    try:
        model = genai.GenerativeModel('gemini-1.5-flash')
        prompt = f"""
        As an expert scheduler, create a schedule based on the provided JSON data.
        You are given a list of tasks and a list of availability slots, which can be on different dates.
        Your job is to assign each task to a time within one of the provided availability slots.
        
        Constraints:
        1.  Prioritize tasks with 'high' priority first.
        2.  Schedule all tasks if possible.
        3.  Do not schedule tasks outside the given availability windows.
        4.  Combine the date from the availability slot with the time to create a full ISO 8601 timestamp for start and end times.
        5. Ensure that the start time of a task is before its end time.
        6.  The `is_daily_routine` flag from the input task must be preserved in the output for that task.
        7. Make sure to give 10 minutes break after the each task time.
        8. Schedule a new breakfast/lunch/dinner break for 30 minutes (even if it not added in the task list), after any task ends if it crosses these common meal times:
            - Breakfast: 7:00 AM – 9:00 AM
            - Lunch: 12:00 PM – 2:00 PM
            - Dinner: 7:00 PM – 9:00 PM
        Input Tasks: {json.dumps([task.dict() for task in schedule_request.tasks])}
        Availability Slots: {json.dumps([avail.dict() for avail in schedule_request.availability])}

        Provide the output as a single, valid JSON object with keys "suggested_schedule" and "notes".
        The "suggested_schedule" key must be an array of objects, each with "task_description", "start_time", "end_time", the original "priority", and the original "is_daily_routine" boolean flag.
        The "notes" key should be a brief, encouraging message for the user, make sure to say the user to stay hydrated after completing each task include the task name in the message. Gently reminds them to eat if any tasks cross these common meal times:
            - Breakfast: 7:00 AM – 8:59 AM
            - Lunch: 12:00 PM – 2:00 PM
            - Dinner: 7:00 PM – 9:00 PM

            Use a warm and supportive tone. Keep it casual but clear.
        Example item in suggested_schedule: {{"task_description": "Finish report", "start_time": "2025-07-21T09:00:00", "end_time": "2025-07-21T10:30:00", "priority": "high", "is_daily_routine": false}}
        """
        response = await model.generate_content_async(prompt)
        cleaned_response_text = response.text.strip().replace("```json", "").replace("```", "")
        ai_data = json.loads(cleaned_response_text)

        sanitized_schedule = []
        for item in ai_data.get("suggested_schedule", []):
            desc = item.get("task_description", "Untitled Task")
            if any(meal in desc.lower() for meal in ["breakfast", "lunch", "dinner"]):
                priority = "high"
            else:
                priority = item.get("priority") if isinstance(item.get("priority"), str) and item.get("priority") else "medium"
            sanitized_schedule.append({
                "task_description": desc,
                "start_time": item.get("start_time"),
                "end_time": item.get("end_time"),
                "priority": priority,
                "is_daily_routine": bool(item.get("is_daily_routine", False))
            })

        return AIScheduleResponse(
            schedule_id=str(uuid.uuid4()),
            suggested_schedule=sanitized_schedule,
            notes=ai_data.get("notes", "Schedule generated successfully!")
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to generate AI schedule: {e}")

@app.post("/api/v1/calendar/sync", tags=["Google Calendar"])
async def sync_to_calendar(sync_request: CalendarSyncRequest, current_user: Dict[str, Any] = Depends(get_current_user)):
    user_email = current_user['email']
    
    try:
        batch_tasks = db.batch()
        tasks_collection_ref = db.collection('tasks')

        for item in sync_request.schedule:
            task_doc_ref = tasks_collection_ref.document()
            batch_tasks.set(task_doc_ref, {
                "description": item.task_description,
                "status": "pending",
                "priority": item.priority,
                "start_time": item.start_time,
                "end_time": item.end_time,
                "is_daily_routine": item.is_daily_routine,
                "owner_email": user_email,
                "created_at": firestore.SERVER_TIMESTAMP,
            })
        
        batch_tasks.commit()
        return {"message": "Schedule successfully saved to the database."}

    except Exception as e:
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=f"An error occurred while saving tasks: {e}")

app.mount("/static", StaticFiles(directory="static"), name="static")

@app.get("/", include_in_schema=False)
async def read_root():
    return FileResponse('static/login.html')

@app.get("/{page_name}.html", include_in_schema=False)
async def read_page(page_name: str):
    file_path = f'static/{page_name}.html'
    if os.path.exists(file_path):
        return FileResponse(file_path)
    raise HTTPException(status_code=404, detail="Page not found")

if __name__ == "__main__":
    uvicorn.run("main:app", host="0.0.0.0", port=8000, reload=True)