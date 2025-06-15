import os
import datetime
import uuid
import json
import io
import webbrowser
from typing import List, Optional, Dict, Any

import uvicorn
from fastapi import FastAPI, Depends, HTTPException, status, Request, Query, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from fastapi.security import OAuth2PasswordBearer
from fastapi.responses import FileResponse, JSONResponse, HTMLResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field
from jose import JWTError, jwt
from dotenv import load_dotenv
from google.oauth2 import id_token as google_id_token
from google.auth.transport import requests as google_requests
from google.oauth2.credentials import Credentials
from google_auth_oauthlib.flow import Flow
from googleapiclient.discovery import build
import google.generativeai as genai

import firebase_admin
from firebase_admin import credentials, firestore

import openpyxl
import docx
from pypdf import PdfReader

from pywebpush import webpush, WebPushException
from apscheduler.schedulers.asyncio import AsyncIOScheduler
from apscheduler.triggers.date import DateTrigger
from apscheduler.triggers.cron import CronTrigger

import traceback

load_dotenv()

VAPID_PRIVATE_KEY = os.getenv("VAPID_PRIVATE_KEY")
VAPID_PUBLIC_KEY = os.getenv("VAPID_PUBLIC_KEY")
VAPID_CLAIMS = {"sub": "mailto:your-email@example.com"}

scheduler = AsyncIOScheduler()

GOOGLE_CLIENT_ID = os.getenv("GOOGLE_CLIENT_ID")
JWT_SECRET_KEY = os.getenv("JWT_SECRET_KEY")
JWT_ALGORITHM = os.getenv("JWT_ALGORITHM")
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")
FIREBASE_SERVICE_ACCOUNT_KEY_PATH = os.getenv("FIREBASE_SERVICE_ACCOUNT_KEY_PATH")

if GEMINI_API_KEY:
    genai.configure(api_key=GEMINI_API_KEY)
else:
    print("Warning: GEMINI_API_KEY not found. AI features will be disabled.")

if not FIREBASE_SERVICE_ACCOUNT_KEY_PATH or not os.path.exists(FIREBASE_SERVICE_ACCOUNT_KEY_PATH):
    raise ValueError("Firebase service account key path not found or invalid.")
cred = credentials.Certificate(FIREBASE_SERVICE_ACCOUNT_KEY_PATH)
firebase_admin.initialize_app(cred)
db = firestore.client()

app = FastAPI(title="SnapTask API with Firebase")


@app.on_event("startup")
def on_startup():
    scheduler.start()
    scheduler.add_job(reset_daily_routine_tasks, CronTrigger(hour=0, minute=0), id="reset_daily_routine_tasks", replace_existing=True)
    print("APScheduler started.")


@app.on_event("shutdown")
def on_shutdown():
    scheduler.shutdown()


@app.middleware("http")
async def add_security_headers(request: Request, call_next):
    response = await call_next(request)
    response.headers["Cross-Origin-Opener-Policy"] = "same-origin-allow-popups"
    response.headers["Cross-Origin-Embedder-Policy"] = "require-corp"
    return response


app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5500", "http://127.0.0.1:5500", "http://localhost:8000", "http://127.0.0.1:8000"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

oauth2_scheme = OAuth2PasswordBearer(tokenUrl="/auth/google/callback")


class PushSubscription(BaseModel):
    endpoint: str
    keys: Dict[str, str]


class TaskBase(BaseModel):
    description: str
    priority: str = "medium"
    estimated_duration_minutes: Optional[int] = 30
    start_time: Optional[str] = None
    end_time: Optional[str] = None
    is_daily_routine: bool = False
    updated_at: Optional[str] = None


class TaskCreate(TaskBase):
    pass


class TaskSchema(TaskBase):
    id: str
    status: str
    owner_email: str


class Availability(BaseModel):
    date: str
    start_time: str
    end_time: str


class ScheduleRequest(BaseModel):
    tasks: List[TaskCreate]
    availability: List[Availability]


class ScheduledTask(BaseModel):
    task_description: str
    start_time: str
    end_time: str
    priority: str
    is_daily_routine: bool


class AIScheduleResponse(BaseModel):
    schedule_id: str
    suggested_schedule: List[ScheduledTask]
    notes: str


class CalendarSyncRequest(BaseModel):
    schedule: List[ScheduledTask]


class OcrResponse(BaseModel):
    tasks: List[str]


class Subtask(BaseModel):
    description: str
    duration_minutes: Optional[int] = None


class BreakdownResponse(BaseModel):
    subtasks: List[Subtask]


def create_access_token(data: dict, expires_delta: Optional[datetime.timedelta] = None):
    to_encode = data.copy()
    expire = datetime.datetime.now(datetime.timezone.utc) + (expires_delta or datetime.timedelta(days=1))
    to_encode.update({"exp": expire})
    return jwt.encode(to_encode, JWT_SECRET_KEY, algorithm=JWT_ALGORITHM)


async def get_current_user(token: str = Depends(oauth2_scheme)) -> Dict[str, Any]:
    credentials_exception = HTTPException(
        status_code=status.HTTP_401_UNAUTHORIZED,
        detail="Could not validate credentials",
        headers={"WWW-Authenticate": "Bearer"},
    )
    try:
        payload = jwt.decode(token, JWT_SECRET_KEY, algorithms=[JWT_ALGORITHM])
        email: str = payload.get("email")
        if email is None:
            raise credentials_exception
    except JWTError:
        raise credentials_exception

    user_doc_ref = db.collection('users').document(email)
    user_doc = user_doc_ref.get()
    if not user_doc.exists:
        raise credentials_exception

    user_data = user_doc.to_dict()
    user_data['email'] = email
    return user_data


def send_push_notification(subscription_info: dict, title: str, body: str, data: dict = None):
    try:
        payload = {"title": title, "body": body, "url": "/static/dashboard.html"}
        if data and "url" in data:
            payload["url"] = data["url"]

        webpush(
            subscription_info=subscription_info,
            data=json.dumps(payload),
            vapid_private_key=VAPID_PRIVATE_KEY,
            vapid_claims=VAPID_CLAIMS.copy()
        )
        print(f"Push notification sent successfully for: {title}")
    except WebPushException as ex:
        if ex.response.status_code == 410:
             print("Subscription has expired or is no longer valid.")
        else:
            print(f"Web push failed: {ex}")
    except Exception as e:
        print(f"An error occurred in send_push_notification: {e}")


def reset_daily_routine_tasks():
    try:
        users = db.collection('users').stream()
        for user in users:
            user_email = user.id
            tasks_to_reset_query = db.collection('tasks').where('owner_email', '==', user_email).where('is_daily_routine', '==', True)
            batch = db.batch()
            for task in tasks_to_reset_query.stream():
                task_data = task.to_dict()
                if task_data.get('is_daily_routine', False):
                    batch.update(task.reference, {
                        'status': 'pending',
                        'updated_at': datetime.datetime.now(datetime.timezone.utc).isoformat()
                    })
            batch.commit()
        print("Daily routine tasks reset at midnight.")
    except Exception as e:
        print(f"Error resetting daily routine tasks: {e}")


@app.get("/vapid_public_key", tags=["Push Notifications"])
def get_vapid_public_key(current_user: Dict[str, Any] = Depends(get_current_user)):
    return {"public_key": VAPID_PUBLIC_KEY}


@app.post("/push/subscribe", status_code=201, tags=["Push Notifications"])
async def subscribe_to_push(subscription: PushSubscription, current_user: Dict[str, Any] = Depends(get_current_user)):
    user_email = current_user['email']
    subscription_dict = subscription.dict()
    sub_collection_ref = db.collection('users').document(user_email).collection('push_subscriptions')
    endpoint_hash = str(uuid.uuid5(uuid.NAMESPACE_URL, subscription.endpoint))
    sub_collection_ref.document(endpoint_hash).set(subscription_dict)
    return {"message": "Subscription saved successfully"}


@app.post("/auth/google/callback", tags=["Authentication"])
async def google_auth_callback(google_token: dict):
    try:
        if not GOOGLE_CLIENT_ID:
            raise HTTPException(status_code=500, detail="GOOGLE_CLIENT_ID is not set on the backend.")
        idinfo = google_id_token.verify_oauth2_token(google_token['token'], google_requests.Request(), GOOGLE_CLIENT_ID)
        email = idinfo['email']
        user_doc_ref = db.collection('users').document(email)
        user_doc = user_doc_ref.get()
        if not user_doc.exists:
            user_data = {
                "email": email,
                "name": idinfo.get('name', 'New User'),
                "picture_url": idinfo.get('picture'),
                "created_at": firestore.SERVER_TIMESTAMP,
                "calendar_credentials": None,
                "last_daily_refresh": None
            }
            user_doc_ref.set(user_data)
        app_token = create_access_token(data={"email": email})
        return {
            "app_token": app_token,
            "user_info": {
                "name": idinfo.get('name'),
                "email": email,
                "picture_url": idinfo.get('picture')
            }
        }
    except ValueError as e:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail=f"Invalid Google token: {e}")
    except Exception as e:
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=f"Internal error: {e}")


@app.get("/tasks/", response_model=List[TaskSchema], tags=["Tasks"])
async def read_tasks(current_user: Dict[str, Any] = Depends(get_current_user)):
    user_email = current_user['email']
    today_str = datetime.date.today().isoformat()
    last_refresh_str = current_user.get('last_daily_refresh')

    if last_refresh_str != today_str:
        print(f"Running daily task refresh for {user_email}...")
        tasks_to_reset_query = db.collection('tasks').where('owner_email', '==', user_email).where('is_daily_routine', '==', True).where('status', '==', 'completed')
        batch = db.batch()
        for task in tasks_to_reset_query.stream():
            batch.update(task.reference, {
                'status': 'pending',
                'updated_at': datetime.datetime.now(datetime.timezone.utc).isoformat()
            })
        batch.commit()
        db.collection('users').document(user_email).update({'last_daily_refresh': today_str})
        print("Daily task refresh complete.")

    tasks_stream = db.collection('tasks').where('owner_email', '==', user_email).stream()
    tasks_list = []
    for task in tasks_stream:
        task_data = task.to_dict()
        task_data['id'] = task.id
        tasks_list.append(task_data)
    return tasks_list


@app.patch("/tasks/{task_id}/complete", status_code=status.HTTP_200_OK, tags=["Tasks"])
async def complete_task(task_id: str, current_user: Dict[str, Any] = Depends(get_current_user)):
    user_email = current_user['email']
    task_ref = db.collection('tasks').document(task_id)
    task_doc = task_ref.get()

    if not task_doc.exists:
        raise HTTPException(status_code=404, detail="Task not found")
    
    if task_doc.to_dict().get('owner_email') != user_email:
        raise HTTPException(status_code=403, detail="Not authorized to modify this task")

    task_ref.update({
        "status": "completed",
        "updated_at": datetime.datetime.now(datetime.timezone.utc).isoformat()
    })
    return {"message": "Task marked as completed."}


@app.delete("/tasks/{task_id}", status_code=status.HTTP_200_OK, tags=["Tasks"])
async def delete_task(task_id: str, current_user: Dict[str, Any] = Depends(get_current_user)):
    user_email = current_user['email']
    task_ref = db.collection('tasks').document(task_id)
    task_doc = task_ref.get()

    if not task_doc.exists:
        raise HTTPException(status_code=404, detail="Task not found")

    if task_doc.to_dict().get('owner_email') != user_email:
        raise HTTPException(status_code=403, detail="Not authorized to delete this task")

    task_ref.delete()
    return {"message": "Task deleted successfully."}


@app.delete("/tasks/reset", status_code=status.HTTP_204_NO_CONTENT, tags=["Tasks"])
async def reset_tasks(current_user: Dict[str, Any] = Depends(get_current_user)):
    user_email = current_user['email']
    tasks_query = db.collection('tasks').where('owner_email', '==', user_email)
    tasks_docs = tasks_query.stream()

    batch = db.batch()
    doc_count = 0
    for doc in tasks_docs:
        batch.delete(doc.reference)
        doc_count += 1
    
    if doc_count > 0:
        batch.commit()
    
    return {"message": f"{doc_count} tasks have been reset."}


@app.post("/tasks/{task_id}/breakdown", response_model=BreakdownResponse, tags=["AI Processing"])
async def breakdown_task_with_ai(task_id: str, current_user: Dict[str, Any] = Depends(get_current_user)):
    user_email = current_user['email']
    task_ref = db.collection('tasks').document(task_id)
    task_doc = task_ref.get()

    if not task_doc.exists or task_doc.to_dict().get('owner_email') != user_email:
        raise HTTPException(status_code=404, detail="Task not found or not authorized")

    task_description = task_doc.to_dict().get('description')
    start_time = task_doc.to_dict().get('start_time')
    end_time = task_doc.to_dict().get('end_time')
    estimated_duration = 120
    if start_time and end_time:
        try:
            start_dt = datetime.datetime.fromisoformat(start_time)
            end_dt = datetime.datetime.fromisoformat(end_time)
            estimated_duration = int((end_dt - start_dt).total_seconds() // 60)
        except Exception:
            try:
                estimated_duration = int(end_time) - int(start_time)
            except Exception:
                estimated_duration = 120
    else:
        estimated_duration = task_doc.to_dict().get('estimated_duration_minutes', 120)
        
    if not task_description:
        raise HTTPException(status_code=400, detail="Task has no description to break down.")
        
    if not GEMINI_API_KEY or not genai:
        raise HTTPException(status_code=503, detail="AI Service is not configured or available.")

    try:
        model = genai.GenerativeModel('gemini-1.5-flash')
        prompt = f'''
        You are a project manager expert. A user has a high-level task and needs help breaking it down into smaller, actionable steps.
        The main task is: "{task_description}"
        The total estimated time for this task is {estimated_duration} minutes.
        Generate a list of 3 to 5 clear, concise sub-tasks that would help complete this main task.
        For each sub-task, suggest a reasonable estimated duration in minutes (as an integer) so that the sum of all sub-task durations is close to the total estimated time.
        In the description of each sub-task, append the timing in parentheses, e.g., "Set a budget and create a guest list (20 min)".
        Return the result as a single, valid JSON object with one key: "subtasks".
        The value of "subtasks" should be an array of objects, where each object has two keys: "description" (string, with timing in parentheses) and "duration_minutes" (integer).
        Example:
        If the main task is "Plan a birthday party" and the total estimated time is 120 minutes, your output should look like this:
        {{
            "subtasks": [
                {{"description": "Set a budget and create a guest list (20 min)", "duration_minutes": 20}},
                {{"description": "Choose a date, time, and venue (25 min)", "duration_minutes": 25}},
                {{"description": "Send out invitations and track RSVPs (25 min)", "duration_minutes": 25}},
                {{"description": "Plan the menu and order a cake (25 min)", "duration_minutes": 25}},
                {{"description": "Organize decorations and entertainment (25 min)", "duration_minutes": 25}}
            ]
        }}
        '''
        response = await model.generate_content_async(prompt)
        cleaned_response_text = response.text.strip().replace("```json", "").replace("```", "")
        ai_data = json.loads(cleaned_response_text)
        subtasks = ai_data.get("subtasks", [])
        total_ai_minutes = sum(sub.get("duration_minutes", 0) for sub in subtasks if isinstance(sub.get("duration_minutes"), int))
        if subtasks and total_ai_minutes != estimated_duration:
            import re
            scaled = []
            running_total = 0
            for i, sub in enumerate(subtasks):
                if i == len(subtasks) - 1:
                    new_minutes = estimated_duration - running_total
                else:
                    ratio = sub.get("duration_minutes", 0) / total_ai_minutes if total_ai_minutes else 0
                    new_minutes = max(1, round(ratio * estimated_duration))
                    running_total += new_minutes
                sub["duration_minutes"] = new_minutes
                sub["description"] = re.sub(r"\(\d+\s*min\)", f"({new_minutes} min)", sub["description"])
            for sub in subtasks:
                try:
                    sub["duration_minutes"] = int(sub["duration_minutes"])
                except Exception:
                    sub["duration_minutes"] = None
        return BreakdownResponse(subtasks=subtasks)
    except Exception as e:
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=f"Failed to process task breakdown: {e}")


@app.post("/process/ai/generate_schedule", response_model=AIScheduleResponse, tags=["AI Processing"])
async def generate_ai_schedule(schedule_request: ScheduleRequest, current_user: Dict[str, Any] = Depends(get_current_user)):
    if not GEMINI_API_KEY or not genai:
        raise HTTPException(status_code=503, detail="AI Service is not configured or available.")
    try:
        model = genai.GenerativeModel('gemini-1.5-flash')
        prompt = f"""
        As an expert scheduler, create a schedule based on the provided JSON data.
        You are given a list of tasks and a list of availability slots, which can be on different dates.
        Your job is to assign each task to a time within one of the provided availability slots.
        
        Constraints:
        1.  Prioritize tasks with 'high' priority first.
        2.  Schedule all tasks if possible.
        3.  Do not schedule tasks outside the given availability windows.
        4.  Combine the date from the availability slot with the time to create a full ISO 8601 timestamp for start and end times.
        5.  Ensure that the start time of a task is before its end time.
        6.  The `is_daily_routine` flag from the input task must be preserved in the output for that task.
        7.  Make sure to give 10 minutes break after the each task time.
        8.  Schedule a new breakfast/lunch/dinner break for 30 minutes (even if it not added in the task list), after any task ends if it crosses these common meal times:
            - Breakfast: 7:00 AM – 9:00 AM
            - Lunch: 12:00 PM – 2:00 PM
            - Dinner: 7:00 PM – 9:00 PM
        Input Tasks: {json.dumps([task.dict() for task in schedule_request.tasks])}
        Availability Slots: {json.dumps([avail.dict() for avail in schedule_request.availability])}

        Provide the output as a single, valid JSON object with keys "suggested_schedule" and "notes".
        The "suggested_schedule" key must be an array of objects, each with "task_description", "start_time", "end_time", the original "priority", and the original "is_daily_routine" boolean flag.
        The "notes" key should be a brief, encouraging message for the user, make sure to say the user to stay hydrated after completing each task include the task name in the message. Gently reminds them to eat if any tasks cross these common meal times:
            - Breakfast: 7:00 AM – 8:59 AM
            - Lunch: 12:00 PM – 2:00 PM
            - Dinner: 7:00 PM – 9:00 PM

            Use a warm and supportive tone. Keep it casual but clear.
        Example item in suggested_schedule: {{"task_description": "Finish report", "start_time": "2025-07-21T09:00:00", "end_time": "2025-07-21T10:30:00", "priority": "high", "is_daily_routine": false}}
        """
        response = await model.generate_content_async(prompt)
        cleaned_response_text = response.text.strip().replace("```json", "").replace("```", "")
        ai_data = json.loads(cleaned_response_text)

        sanitized_schedule = []
        for item in ai_data.get("suggested_schedule", []):
            desc = item.get("task_description", "Untitled Task")
            if any(meal in desc.lower() for meal in ["breakfast", "lunch", "dinner"]):
                priority = item.get("priority") if isinstance(item.get("priority"), str) and item.get("priority") else "medium"
                is_daily_routine = True
            else:
                priority = item.get("priority") if isinstance(item.get("priority"), str) and item.get("priority") else "medium"
                is_daily_routine = bool(item.get("is_daily_routine", False))
            sanitized_schedule.append({
                "task_description": desc,
                "start_time": item.get("start_time"),
                "end_time": item.get("end_time"),
                "priority": priority,
                "is_daily_routine": is_daily_routine
            })
        
        return AIScheduleResponse(
            schedule_id=str(uuid.uuid4()),
            suggested_schedule=sanitized_schedule,
            notes=ai_data.get("notes", "Schedule generated successfully!")
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to generate AI schedule: {e}")


@app.post("/api/v1/calendar/sync", tags=["Google Calendar"])
async def sync_to_calendar(sync_request: CalendarSyncRequest, current_user: Dict[str, Any] = Depends(get_current_user)):
    user_email = current_user['email']
    
    try:
        batch_tasks = db.batch()
        tasks_collection_ref = db.collection('tasks')

        for item in sync_request.schedule:
            task_doc_ref = tasks_collection_ref.document()
            batch_tasks.set(task_doc_ref, {
                "description": item.task_description,
                "status": "pending",
                "priority": item.priority,
                "start_time": item.start_time,
                "end_time": item.end_time,
                "is_daily_routine": item.is_daily_routine,
                "owner_email": user_email,
                "created_at": firestore.SERVER_TIMESTAMP,
            })
            
            if item.start_time:
                try:
                    start_time_dt = datetime.datetime.fromisoformat(item.start_time.replace("Z", "+00:00"))
                    notification_time = start_time_dt - datetime.timedelta(minutes=10)
                    now = datetime.datetime.now(datetime.timezone.utc)

                    if notification_time > now:
                        subs_ref = db.collection('users').document(user_email).collection('push_subscriptions').stream()
                        subscriptions = [sub.to_dict() for sub in subs_ref]
                        
                        for sub_info in subscriptions:
                            job_id = f"push_{user_email}_{task_doc_ref.id}_{sub_info['keys']['p256dh']}"
                            scheduler.add_job(
                                send_push_notification,
                                trigger=DateTrigger(run_date=notification_time),
                                args=[
                                    sub_info,
                                    f"Task starting: {item.task_description}",
                                    "This task is scheduled to begin in 10 minutes. Get ready!",
                                    # Add the URL to open on click
                                    {"url": "/static/dashboard.html"}
                                ],
                                id=job_id,
                                replace_existing=True
                            )
                            print(f"Scheduled notification for job {job_id} at {notification_time}")

                except Exception as e:
                    print(f"Error scheduling notification for task '{item.task_description}': {e}")


        batch_tasks.commit()
        return {"message": "Schedule successfully saved and notifications scheduled."}

    except Exception as e:
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=f"An error occurred while saving tasks: {e}")


@app.get("/api/v1/dashboard/summary", tags=["Dashboard"])
async def get_dashboard_summary(current_user: Dict[str, Any] = Depends(get_current_user)):
    user_email = current_user['email']
    tasks_ref = db.collection('tasks').where('owner_email', '==', user_email).stream()
    total_tasks = 0
    completed_tasks = 0
    for task in tasks_ref:
        total_tasks += 1
        if task.to_dict().get('status') == 'completed':
            completed_tasks += 1
    return {"totalTasks": total_tasks, "completedTasks": completed_tasks}

app.mount("/", StaticFiles(directory="static", html=True), name="static-root")


@app.get("/", include_in_schema=False)
async def read_root():
    return FileResponse('index.html')


@app.get("/{page_name}.html", include_in_schema=False)
async def read_page(page_name: str):
    file_path = f'static/{page_name}.html'
    if os.path.exists(file_path):
        return FileResponse(file_path)
    raise HTTPException(status_code=404, detail="Page not found")


@app.post("/process/file/extract_tasks", tags=["AI Processing"])
async def extract_tasks_from_file(
    file: UploadFile = File(...),
    current_user: Dict[str, Any] = Depends(get_current_user)
):
    if not GEMINI_API_KEY or not genai:
        raise HTTPException(status_code=503, detail="AI Service is not configured or available.")

    content_type = file.content_type
    file_content = await file.read()
    extracted_text = ""

    try:
        if "image" in content_type:
            model = genai.GenerativeModel('gemini-1.5-flash')
            image_parts = [{"mime_type": content_type, "data": file_content}]
            prompt = (
                "Analyze the image and extract all distinct tasks or to-do list items. "
                "For each task, if a priority (high/medium/low) or estimated duration (in minutes) is mentioned, extract those as well. "
                "Return a JSON object with a 'tasks' key containing an array of objects, each with 'description', 'priority', and 'estimated_duration_minutes' if available."
            )
            response = await model.generate_content_async([prompt, *image_parts])
            extracted_text = response.text
        elif content_type == 'application/pdf':
            reader = PdfReader(io.BytesIO(file_content))
            for page in reader.pages:
                extracted_text += page.extract_text() + "\n"
        elif content_type == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document':
            doc = docx.Document(io.BytesIO(file_content))
            for para in doc.paragraphs:
                extracted_text += para.text + "\n"
        elif content_type == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet':
            workbook = openpyxl.load_workbook(io.BytesIO(file_content))
            sheet = workbook.active
            for row in sheet.iter_rows(values_only=True):
                extracted_text += " ".join([str(cell) for cell in row if cell is not None]) + "\n"
        else:
            raise HTTPException(status_code=400, detail=f"Unsupported file type: {content_type}")

        if "image" not in content_type:
            model = genai.GenerativeModel('gemini-1.5-flash')
            prompt = f"""
            Analyze the following text content extracted from a document.\n
            Identify all distinct tasks or to-do list items from this text.\n
            For each task, if a priority (high/medium/low) or estimated duration (in minutes) is mentioned, extract those as well.\n
            Return the result as a single, valid JSON object with one key: 'tasks'.\n
            The value of 'tasks' should be an array of objects, each with:\n
            - 'description': the task description (string)\n            - 'priority': the priority if found (string: high, medium, or low; default to medium if not found)\n            - 'estimated_duration_minutes': integer if found, else 30\n
            If no tasks are found, return an empty array.\n
            Here is the text:\n---\n{extracted_text}\n---\n"""
            response = await model.generate_content_async(prompt)
            cleaned_response_text = response.text.strip().replace("```json", "").replace("```", "")
        else:
            cleaned_response_text = extracted_text.strip().replace("```json", "").replace("```", "")

        ai_data = json.loads(cleaned_response_text)
        tasks = []
        for t in ai_data.get("tasks", []):
            if isinstance(t, dict):
                desc = t.get("description") or t.get("task") or "Untitled Task"
                priority = t.get("priority", "medium")
                try:
                    duration = int(t.get("estimated_duration_minutes", 30))
                except Exception:
                    duration = 30
                tasks.append({
                    "description": desc,
                    "priority": priority,
                    "estimated_duration_minutes": duration
                })
            elif isinstance(t, str):
                tasks.append({
                    "description": t,
                    "priority": "medium",
                    "estimated_duration_minutes": 30
                })
        return {"tasks": tasks}
    except Exception as e:
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=f"Failed to process file: {str(e)}")
